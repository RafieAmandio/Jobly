import io

import pytest
from docx import Document

from jobly.services.cv_parser import extract_text_from_docx, extract_text_from_pdf


def test_extract_text_invalid_pdf():
    with pytest.raises(Exception):
        extract_text_from_pdf(b"not a pdf")


def test_extract_text_from_docx():
    doc = Document()
    doc.add_paragraph("John Doe")
    doc.add_paragraph("Software Engineer")
    buffer = io.BytesIO()
    doc.save(buffer)

    text = extract_text_from_docx(buffer.getvalue())

    assert "John Doe" in text
    assert "Software Engineer" in text
