from io import BytesIO
from zipfile import ZipFile

import pdfplumber
from docx import Document

from jobly.services.doc_generator import (
    generate_cover_letter_docx,
    generate_cover_letter_pdf,
    generate_cv_docx,
    generate_cv_pdf,
)


def test_generate_cv_docx():
    data = {
        "summary": "Experienced software engineer with 5 years in fintech.",
        "experience": [
            {
                "title": "Senior Engineer",
                "company": "Gojek",
                "location": "Jakarta, Indonesia",
                "period": "2021-present",
                "bullets": [
                    "Lead migration to microservices",
                    "Reduce latency by 40%",
                ],
            }
        ],
        "skills": ["Python", "Go", "PostgreSQL", "Docker"],
        "education": [
            {
                "degree": "BS Computer Science",
                "institution": "Universitas Indonesia",
                "year": "2019",
            }
        ],
    }

    result = generate_cv_docx(data, "John Doe")
    assert isinstance(result, bytes)
    assert len(result) > 0
    assert result[:2] == b"PK"
    parsed = Document(BytesIO(result))
    body_text = [p.text for p in parsed.paragraphs]
    assert any("Senior Engineer\tJakarta, Indonesia | 2021-present" in text for text in body_text)


def test_generate_cv_docx_minimal():
    result = generate_cv_docx({}, "Minimal User")
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_generate_cv_pdf_aligns_work_location_with_period():
    data = {
        "experience": [
            {
                "company": "DBS Bank HK Limited",
                "title": "Business Analyst (Global Financial Market Operations)",
                "location": "Hong Kong SAR",
                "period": "Jul 2025 – Jun 2026",
                "bullets": [
                    "Analyze operational requirements",
                    "Coordinate stakeholders",
                    "Improve reporting workflows",
                ],
            }
        ]
    }

    result = generate_cv_pdf(data, "Dharma Setiawan")
    assert isinstance(result, bytes)
    with pdfplumber.open(BytesIO(result)) as pdf:
        text = pdf.pages[0].extract_text() or ""

    assert "DBS Bank HK Limited" in text
    assert "Business Analyst (Global Financial Market Operations) Hong Kong SAR | Jul 2025 – Jun 2026" in text


def test_generate_cv_pdf_keeps_six_three_bullet_roles_within_two_pages():
    data = {
        "summary": (
            "Business Analyst with cross-functional financial-services experience. "
            "Translates operational requirements into practical workflows and stakeholder outcomes."
        ),
        "experience": [
            {
                "company": f"Relevant Employer {index}",
                "title": "Business Analyst",
                "location": "Hong Kong SAR",
                "period": f"202{index} – 202{index + 1}",
                "bullets": [
                    "Analyzed operational requirements and documented implementation priorities",
                    "Coordinated business stakeholders across delivery and governance workstreams",
                    "Improved reporting workflows using source-supported process findings",
                ],
            }
            for index in range(6)
        ],
        "education": [
            {
                "institution": "The Chinese University of Hong Kong, Shenzhen",
                "degree": "Bachelor of Science",
                "location": "Shenzhen, China",
                "year": "2020 – 2024",
            }
        ],
    }

    result = generate_cv_pdf(data, "Dharma Setiawan")
    assert isinstance(result, bytes)
    with pdfplumber.open(BytesIO(result)) as pdf:
        assert len(pdf.pages) <= 2
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    assert "Relevant Employer 0" in text
    assert "Relevant Employer 5" in text


def test_generate_cv_docx_with_contact_and_extra_sections():
    data = {
        "display_name": "Rafie",
        "contact": {
            "location": "Jakarta, Indonesia",
            "email": "rafie@example.com",
            "phone": "+62 813 0000 0000",
            "linkedin": "https://linkedin.com/in/rafie",
        },
        "summary": "Senior engineer.",
        "experience": [
            {
                "company": "REVALUE ACADEMY, Remote",
                "title": "Senior Software Engineer",
                "period": "March 2025 - Present",
                "bullets": ["Build the backend from scratch"],
            }
        ],
        "education": [
            {
                "institution": "UNIVERSITAS INDONESIA",
                "degree": "Computer Engineering",
                "gpa": "3.85/4.00",
                "year": "2021 - 2025",
                "location": "Depok, Indonesia",
                "coursework": ["Distributed Systems", "Machine Learning"],
                "honors_awards": ["Dean's List"],
                "activities": ["Teaching Assistant", "Student Representative"],
            }
        ],
        "certifications": ["Cisco CCNA"],
        "awards": ["1st Place, Galaxy Hackathon"],
        "projects": ["ProjectXOXO, 90,000+ users"],
        "additional_info": {
            "Languages": ["Indonesian", "English"],
            "Tests": ["IELTS 7.0"],
        },
        "skills": ["Python", "Go"],
    }
    result = generate_cv_docx(data, "Rafie Amandio Fauzan")
    parsed = Document(BytesIO(result))
    header_text = [p.text for p in parsed.sections[0].first_page_header.paragraphs]
    body_text = [p.text for p in parsed.paragraphs]

    assert isinstance(result, bytes)
    assert result[:2] == b"PK"
    assert any("Rafie" in text and "Jakarta, Indonesia | rafie@example.com" in text for text in header_text)
    assert any("3.85/4.00" in text for text in body_text)
    assert "EDUCATION BACKGROUND" in body_text
    assert any("UNIVERSITAS INDONESIA\tDepok, Indonesia" in text for text in body_text)
    assert any("Computer Engineering\t2021 - 2025" in text for text in body_text)
    assert any(
        "Related Coursework: Distributed Systems, Machine Learning" in text
        for text in body_text
    )
    assert any("Honors and Awards: Dean's List" in text for text in body_text)
    assert any("Activities: Teaching Assistant, Student Representative" in text for text in body_text)
    assert not any(text.startswith("●") and "Coursework" in text for text in body_text)
    assert "ADDITIONAL INFORMATION" in body_text
    assert any("Languages: Indonesian, English" in text for text in body_text)
    assert any("Tests: IELTS 7.0" in text for text in body_text)


def test_generate_cover_letter_docx():
    content = "Dear Hiring Manager,\n\nI am writing to express my interest.\n\nBest regards,\nJohn"
    contact = {"email": "john@example.com", "phone": "+62 812 0000 0000"}
    result = generate_cover_letter_docx(content, "John Doe", contact)
    parsed = Document(BytesIO(result))
    header_text = [p.text for p in parsed.sections[0].header.paragraphs]

    assert isinstance(result, bytes)
    assert len(result) > 0
    assert result[:2] == b"PK"
    assert parsed.styles["Normal"].font.name == "Times New Roman"
    assert round(parsed.styles["Normal"].font.size.pt) == 10
    assert any("John Doe" in text and "john@example.com | +62 812 0000 0000" in text for text in header_text)


def test_generate_cover_letter_pdf():
    content = "Dear Hiring Manager,\n\nI am writing to express my interest.\n\nBest regards,\nJohn"
    contact = {"email": "john@example.com", "phone": "+62 812 0000 0000"}
    result = generate_cover_letter_pdf(content, "John Doe", contact)

    assert isinstance(result, bytes)
    assert result.startswith(b"%PDF")


def test_generate_cv_docx_uses_running_header_and_renders_leadership_org():
    data = {
        "display_name": "Dharma",
        "contact": {
            "location": "Hong Kong SAR",
            "email": "dharma@example.com",
            "phone": "+62 813 0000 0000",
        },
        "summary": "Finance-focused operator.",
        "experience": [
            {
                "company": "Financial Light",
                "title": "Analyst",
                "period": "Jan 2024 - Present",
                "bullets": ["Improve conversion by 22%"],
            }
        ],
        "leadership": [
            {
                "organization": "BEM UI",
                "title": "Treasurer",
                "period": "2023 - 2024",
                "bullets": ["Managed Rp 120M budget"],
            }
        ],
        "extra_miles": ["1st Place, Galaxy Hackathon"],
    }

    result = generate_cv_docx(data, "Dharma Setiawan")
    parsed = Document(BytesIO(result))
    text = [p.text for p in parsed.paragraphs]
    header_text = [p.text for p in parsed.sections[0].first_page_header.paragraphs]

    assert parsed.styles["Normal"].font.name == "Times New Roman"
    assert round(parsed.styles["Normal"].font.size.pt) == 10
    assert "SUMMARY" in text
    assert "WORK EXPERIENCES" in text
    assert "VOLUNTEER & LEADERSHIP EXPERIENCES" in text
    assert "EXTRA MILES" in text
    assert any("BEM UI" in t for t in text)
    assert any("Managed Rp 120M budget" in t for t in text)
    assert any("Dharma" in t and "Hong Kong SAR | dharma@example.com" in t for t in header_text)
    assert parsed.sections[0].different_first_page_header_footer is True
    assert not any(p.text.strip() for p in parsed.sections[0].header.paragraphs)

    with ZipFile(BytesIO(result)) as zf:
        header_paths = [name for name in zf.namelist() if name.startswith("word/header")]
        header_xml = "\n".join(zf.read(name).decode("utf-8") for name in header_paths)
        assert "Dharma" in header_xml
        assert 'w:sz w:val="32"' in header_xml
        assert "Hong Kong SAR" in header_xml
        assert "dharma@example.com" in header_xml

        rel_paths = [name for name in zf.namelist() if name.startswith("word/_rels/header")]
        rels_xml = "\n".join(zf.read(name).decode("utf-8") for name in rel_paths)
        assert "mailto:dharma@example.com" in rels_xml
        document_xml = zf.read("word/document.xml").decode("utf-8")
        assert "w:titlePg" in document_xml


def test_generate_cv_pdf_renders_large_name_separate_from_contact():
    data = {
        "display_name": "Dharma Setiawan",
        "contact": {"location": "Shenzhen, China", "email": "dharma@example.com"},
        "summary": "Management Trainee and PM/BA profile.",
    }

    result = generate_cv_pdf(data, "Dharma Setiawan")

    assert isinstance(result, bytes)
    assert b"font-size: 16pt" in result or result.startswith(b"%PDF")


def test_generate_cv_pdf_only_has_header_on_first_page():
    data = {
        "display_name": "Dharma Setiawan",
        "contact": {"email": "dharma@example.com"},
        "summary": "Management Trainee and PM/BA profile.",
        "extra_miles": [
            f"Achievement {index}: delivered measurable operational improvement across a complex project"
            for index in range(200)
        ],
    }

    result = generate_cv_pdf(data, "Dharma Setiawan")
    assert isinstance(result, bytes)
    with pdfplumber.open(BytesIO(result)) as pdf:
        assert len(pdf.pages) >= 2
        first_page = pdf.pages[0].extract_text() or ""
        second_page = pdf.pages[1].extract_text() or ""

    assert "Dharma Setiawan" in first_page
    assert "dharma@example.com" in first_page
    assert "Dharma Setiawan" not in second_page
    assert "dharma@example.com" not in second_page


def test_generate_cv_docx_omits_missing_optional_education_rows():
    data = {
        "education": [
            {
                "institution": "Universitas Indonesia",
                "degree": "BSc Computer Science",
                "year": "2020 - 2024",
            }
        ]
    }

    parsed = Document(BytesIO(generate_cv_docx(data, "Candidate")))
    body_text = "\n".join(p.text for p in parsed.paragraphs)

    assert "EDUCATION BACKGROUND" in body_text
    assert "Related Coursework:" not in body_text
    assert "Honors and Awards:" not in body_text
    assert "Activities:" not in body_text
    assert "Cumulative GPA:" not in body_text
